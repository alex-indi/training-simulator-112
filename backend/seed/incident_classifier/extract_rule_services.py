"""Связывает только однозначные значения колонки «Главная служба» SRC-006 с каталогом.

Неподтверждённые сокращения остаются без связи и печатаются для сверки.
"""

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from seed.incident_classifier.extract_src006 import SOURCE_NAME

SERVICE_NAMES = {
    "MCHS": "Служба 101",
    "Police": "Служба 102",
    "AMBULANCE": "Служба 103",
    "MOSGAZ": "Служба 104",
    "MOSLIFT": "Мослифт",
    "AUTOROADS": "Автодороги",
    "MOSVODOCANAL": "Мосводоканал",
    "METRO": "Метро",
    "OEK": "ОЭК",
    "MOSGORTRANS": "Мосгортранс",
    "MOESK": "Россети МР",
    "MOEK": "МОЭК",
    "MZD": "МЖД",
    "MGTS": "МГТС",
    "MOSVODOSTOK": "Мосводосток",
    "MOSCOLLECTOR": "Москоллектор",
    "GORMOST": "Гормост",
    "Dep.tszn": "Деп. труда и соц.защиты",
    "ZODD": "ЦОДД",
    "MSPPN": "ГБУ МСППН",
    "DepEco": "Деп. природопользования",
    "ZEMP": "ЦЭМП",
}


def extract(path: Path) -> tuple[list[dict], Counter[str]]:
    if path.name != SOURCE_NAME:
        raise ValueError(f"Ожидается {SOURCE_NAME}")
    services = json.loads(
        (Path(__file__).resolve().parents[1] / "services/services.json").read_text()
    )
    by_name = {row["official_name"].split(" (")[0]: row for row in services}
    if len(by_name) != len(services):
        raise ValueError("Неоднозначные краткие имена в каталоге служб")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    links: list[dict] = []
    unknown: Counter[str] = Counter()
    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if number < 4 or row[4] is None or row[10] is None:
            continue
        raw = row[12]
        if raw is None or not str(raw).strip():
            continue
        for code in (part.strip() for part in str(raw).split(",")):
            prefix = SERVICE_NAMES.get(code)
            if prefix is None:
                unknown[code] += 1
                continue
            service = by_name.get(prefix)
            if service is None:
                raise ValueError(f"В каталоге нет подтверждённой службы {prefix}")
            links.append(
                {
                    "rule_source_reference": f"{SOURCE_NAME}#{sheet.title}!{number}",
                    "service_source_reference": service["source_reference"],
                    "source_reference": f"{SOURCE_NAME}#{sheet.title}!M{number}",
                }
            )
    workbook.close()
    return links, unknown


if __name__ == "__main__":
    result, unresolved = extract(Path(sys.argv[1]))
    output = Path(__file__).with_name("rule_services.json")
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print({"links": len(result), "unresolved": dict(unresolved)})
