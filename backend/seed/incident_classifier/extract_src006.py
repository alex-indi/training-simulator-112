"""Детерминированно преобразует SRC-006 в seed; исходный XLSX не меняет.

Запуск: python -m seed.incident_classifier.extract_src006 <путь-к-SRC-006.xlsx>
Требуется openpyxl. Связи со службами готовятся отдельно после проверки каталога.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

SOURCE_NAME = "SRC-006-incident-classifier-v046-24.xlsx"
FEATURE_COLUMNS = {
    7: ("1", "112 - Признак.1 (тип происшествия)"),
    8: ("2", "112-Признак.2"),
    9: ("3", "112-Признак.3"),
    10: ("additional", "Дополнительные признаки (не влияют на тип происшествия)"),
}


def extract(path: Path) -> dict[str, list[dict]]:
    if path.name != SOURCE_NAME:
        raise ValueError(f"Ожидается {SOURCE_NAME}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rules: list[dict] = []
    features: dict[tuple[str, str, str], dict] = {}
    links: list[dict] = []
    group: str | None = None
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number < 4:
            continue
        code, group_cell, first_feature, incident_type = row[4], row[5], row[6], row[10]
        if first_feature is None and incident_type is None and group_cell:
            group = str(group_cell).strip()
            continue
        if code is None or incident_type is None:
            continue
        if not group:
            raise ValueError(f"Нет группы для строки {row_number}")
        reference = f"{SOURCE_NAME}#{sheet.title}!{row_number}"
        rules.append(
            {
                "source_code": str(code),
                "incident_group": group,
                "final_incident_type": str(incident_type),
                "ekp35_type": str(row[11]) if row[11] is not None else None,
                "source_reference": reference,
            }
        )
        for column, (level, source_column) in FEATURE_COLUMNS.items():
            value = row[column - 1]
            if value is None or not str(value).strip():
                continue
            source_value = str(value)
            key = (level, source_column, source_value)
            features.setdefault(
                key,
                {
                    "name": source_value.strip(),
                    "level": level,
                    "source_column": source_column,
                    "source_value": source_value,
                },
            )
            links.append({"rule_source_reference": reference, "feature_key": list(key)})
    workbook.close()
    if not rules:
        raise ValueError("В SRC-006 не найдены правила")
    return {"rules": rules, "features": list(features.values()), "rule_features": links}


def main() -> None:
    result = extract(Path(sys.argv[1]))
    root = Path(__file__).resolve().parent
    for name, rows in result.items():
        (root / f"{name}.json").write_text(
            json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    print({name: len(rows) for name, rows in result.items()})


if __name__ == "__main__":
    main()
